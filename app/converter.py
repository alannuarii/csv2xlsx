"""CSV to XLSX converter module.

Provides encoding detection, delimiter detection, smart typing,
pratitinjau (preview) data, dan pembuatan workbook Excel berformat profesional.
"""

from __future__ import annotations

import csv
import datetime
import io
import math
import re
import zipfile
from typing import Any, List, Optional, Tuple

import charset_normalizer
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50 MB
ALLOWED_DELIMITERS = [",", ";", "\t", "|"]
ILLEGAL_EXCEL_CHARS_RE = re.compile(r"[\\/?*:[\]]")


class ConversionError(Exception):
    """Exception raised when CSV parsing or conversion fails."""
    pass


def validate_file_content(data: bytes) -> None:
    """Validate that uploaded content is not empty and not an unsupported binary file.

    Raises:
        ConversionError: If file is empty or contains binary null bytes.
    """
    if not data or len(data.strip()) == 0:
        raise ConversionError("File CSV kosong atau tidak memiliki data.")

    # Check for binary null bytes in the first 8KB
    sample = data[:8192]
    if b"\x00" in sample:
        raise ConversionError("File terdeteksi sebagai biner dan bukan teks CSV/TSV yang valid.")


def validate_xlsx_content(data: bytes) -> None:
    """Validate that uploaded content is a valid, readable Excel (.xlsx) file.

    Raises:
        ConversionError: If file is empty or not a valid XLSX archive.
    """
    if not data or len(data.strip()) == 0:
        raise ConversionError("File Excel kosong atau tidak memiliki data.")

    # XLSX is a ZIP package starting with PK\x03\x04
    if not data.startswith(b"PK\x03\x04"):
        raise ConversionError("File bukan merupakan format spreadsheet Excel (.xlsx) yang valid.")

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            namelist = zf.namelist()
            # Standard XLSX files contain workbook.xml or [Content_Types].xml
            if not any("xl/workbook.xml" in name or "[Content_Types].xml" in name for name in namelist):
                raise ConversionError("Struktur file Excel (.xlsx) tidak valid atau rusak.")
    except (zipfile.BadZipFile, Exception):
        raise ConversionError("File Excel (.xlsx) rusak atau tidak dapat dibaca.")


def detect_encoding(raw_bytes: bytes) -> str:
    """Detect text encoding using UTF-8 validation and charset-normalizer fallback.

    Args:
        raw_bytes: The raw bytes of the file.

    Returns:
        The detected encoding name (e.g. 'utf-8', 'utf-8-sig', 'latin-1', 'windows-1252').
    """
    if not raw_bytes:
        return "utf-8"

    # Check for UTF-8 BOM
    if raw_bytes.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"

    sample = raw_bytes[:65536]

    # Quick & precise check: if valid UTF-8, prefer UTF-8
    try:
        sample.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass

    # Use charset_normalizer on non-UTF8 sample
    try:
        results = charset_normalizer.from_bytes(sample)
        best = results.best()
        if best is not None and best.encoding:
            enc = best.encoding.lower()
            if enc in ("utf_8", "utf8"):
                return "utf-8"
            if enc in ("windows-1252", "cp1252"):
                return "windows-1252"
            if enc in ("iso-8859-1", "latin1", "latin-1"):
                return "latin-1"
            return enc
    except Exception:
        pass

    # Fallback cascade for non-UTF-8 bytes
    for candidate in ["windows-1252", "latin-1"]:
        try:
            sample.decode(candidate)
            return candidate
        except (UnicodeDecodeError, LookupError):
            continue

    return "latin-1"


def detect_delimiter(sample_text: str, candidate_delimiters: Optional[List[str]] = None) -> str:
    """Detect CSV delimiter using csv.Sniffer with robust fallback heuristics.

    Args:
        sample_text: Decoded sample lines from the CSV.
        candidate_delimiters: List of delimiters to test.

    Returns:
        The detected delimiter character.
    """
    if candidate_delimiters is None:
        candidate_delimiters = ALLOWED_DELIMITERS

    # Clean empty lines
    lines = [line for line in sample_text.splitlines() if line.strip()]
    if not lines:
        return ","

    test_sample = "\n".join(lines[:30])

    # 1. Try standard Python Sniffer
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(test_sample, delimiters="".join(candidate_delimiters))
        if dialect.delimiter in candidate_delimiters:
            return dialect.delimiter
    except Exception:
        pass

    # 2. Heuristic fallback: count frequency and consistency per line
    best_delim = ","
    best_score = -1.0

    for d in candidate_delimiters:
        counts = [line.count(d) for line in lines[:20]]
        avg_count = sum(counts) / len(counts) if counts else 0
        if avg_count == 0:
            continue

        # Check consistency (how many lines have the same count)
        mode_count = max(set(counts), key=counts.count) if counts else 0
        if mode_count == 0:
            continue

        consistency = counts.count(mode_count) / len(counts)
        # Score favors higher consistency and reasonable delimiter count
        score = consistency * 10 + avg_count
        if score > best_score:
            best_score = score
            best_delim = d

    return best_delim


def parse_cell_value(val: str) -> Any:
    """Convert text to typed values (int, float, bool, or str).

    Preserves leading zeros on numbers (e.g. phone numbers, zip codes, IDs)
    so they are not corrupted when opened in Excel.

    Args:
        val: The raw cell string value.

    Returns:
        Converted value as int, float, bool, or original str.
    """
    if val is None:
        return ""

    s = val.strip()
    if not s:
        return ""

    # Boolean detection
    s_lower = s.lower()
    if s_lower == "true":
        return True
    if s_lower == "false":
        return False

    # Check for integer
    # Avoid converting numbers with leading zeroes (e.g. '0123', '0812345678')
    if (s.isdigit() or (s.startswith("-") and s[1:].isdigit())):
        num_str = s.lstrip("-")
        if len(num_str) > 1 and num_str.startswith("0"):
            # Keep as string to protect leading zeroes
            return s
        try:
            return int(s)
        except ValueError:
            pass

    # Check for floating point
    if "." in s:
        try:
            f_val = float(s)
            if not (math.isnan(f_val) or math.isinf(f_val)):
                return f_val
        except ValueError:
            pass

    return val


def sanitize_sheet_name(name: str) -> str:
    """Sanitize Excel worksheet name according to Excel rules:
    - Max 31 characters
    - Cannot contain: \\ / ? * : [ ]
    - Default to 'Sheet1' if empty.
    """
    if not name or not name.strip():
        return "Sheet1"

    sanitized = ILLEGAL_EXCEL_CHARS_RE.sub("_", name.strip())
    # Excel worksheet names cannot start or end with single quote (')
    sanitized = sanitized.strip("'")
    sanitized = sanitized[:31].strip()

    return sanitized or "Sheet1"


def read_csv_preview(
    file_bytes: bytes,
    manual_delimiter: Optional[str] = None,
    max_preview_rows: int = 20,
    filename: str = "",
) -> dict:
    """Parse CSV bytes to generate summary statistics and preview rows.

    Args:
        file_bytes: Raw bytes of uploaded file.
        manual_delimiter: Delimiter override if specified by user.
        max_preview_rows: Number of preview rows to return (default: 20).
        filename: Original file name.

    Returns:
        Dictionary containing metadata, statistics, headers, and preview rows.
    """
    validate_file_content(file_bytes)

    detected_enc = detect_encoding(file_bytes)
    try:
        text_content = file_bytes.decode(detected_enc)
    except UnicodeDecodeError:
        # Fallback to latin-1 which can decode any byte sequence
        text_content = file_bytes.decode("latin-1", errors="replace")
        detected_enc = "latin-1"

    # Normalize line endings
    text_content = text_content.replace("\r\n", "\n").replace("\r", "\n")

    # Delimiter resolution
    if manual_delimiter and manual_delimiter in ALLOWED_DELIMITERS:
        delim = manual_delimiter
    else:
        delim = detect_delimiter(text_content[:32768])

    reader = csv.reader(io.StringIO(text_content), delimiter=delim)

    rows: List[List[str]] = []
    total_rows = 0
    headers: List[str] = []

    for idx, row in enumerate(reader):
        total_rows += 1
        if idx == 0:
            headers = [str(c).strip() for c in row]
        elif len(rows) < max_preview_rows:
            rows.append([str(c) for c in row])

    # If file was 1 row or headers empty, handle edge cases
    if total_rows == 0:
        raise ConversionError("File CSV tidak mengandung data.")

    total_cols = len(headers) if headers else 0
    if rows:
        total_cols = max(total_cols, max(len(r) for r in rows))

    return {
        "success": True,
        "filename": filename,
        "detected_delimiter": delim,
        "detected_encoding": detected_enc,
        "total_rows": total_rows,
        "total_columns": total_cols,
        "headers": headers,
        "preview_rows": rows,
    }


def convert_csv_to_xlsx(
    file_bytes: bytes,
    output_stream: io.BufferedIOBase,
    sheet_name: Optional[str] = None,
    manual_delimiter: Optional[str] = None,
    apply_autofilter: bool = True,
    apply_autowidth: bool = True,
) -> None:
    """Convert CSV bytes to a styled Excel (.xlsx) file written to output_stream.

    Args:
        file_bytes: Raw bytes of uploaded file.
        output_stream: Writable binary stream (e.g. BytesIO or open temp file).
        sheet_name: Worksheet title.
        manual_delimiter: Delimiter override if specified by user.
        apply_autofilter: Whether to enable Excel auto-filter on header.
        apply_autowidth: Whether to auto-fit column widths proportionally.
    """
    validate_file_content(file_bytes)

    detected_enc = detect_encoding(file_bytes)
    try:
        text_content = file_bytes.decode(detected_enc)
    except UnicodeDecodeError:
        text_content = file_bytes.decode("latin-1", errors="replace")

    text_content = text_content.replace("\r\n", "\n").replace("\r", "\n")

    if manual_delimiter and manual_delimiter in ALLOWED_DELIMITERS:
        delim = manual_delimiter
    else:
        delim = detect_delimiter(text_content[:32768])

    reader = csv.reader(io.StringIO(text_content), delimiter=delim)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_name(sheet_name or "Data")

    # Styling definitions
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    # Modern dark slate blue / charcoal fill for header
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=11)
    data_alignment = Alignment(vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    header_border = Border(
        left=Side(style="thin", color="334155"),
        right=Side(style="thin", color="334155"),
        top=Side(style="thin", color="334155"),
        bottom=Side(style="medium", color="0F172A"),
    )

    max_lens: dict[int, int] = {}
    row_count = 0
    col_count = 0

    for r_idx, row in enumerate(reader, start=1):
        row_count += 1
        col_count = max(col_count, len(row))

        for c_idx, raw_cell in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)

            if r_idx == 1:
                # Header row
                cell.value = str(raw_cell).strip()
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = header_border
                val_str = str(cell.value)
            else:
                # Data row
                typed_val = parse_cell_value(raw_cell)
                cell.value = typed_val
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                val_str = str(typed_val) if typed_val is not None else ""

            # Track column length for auto-width
            val_len = len(val_str)
            if c_idx not in max_lens or val_len > max_lens[c_idx]:
                max_lens[c_idx] = val_len

    if row_count == 0:
        raise ConversionError("Tidak ada baris data yang ditemukan dalam file CSV.")

    # Freeze header row
    ws.freeze_panes = "A2"

    # Apply Auto-Filter on header
    if apply_autofilter and row_count >= 1 and col_count >= 1:
        start_letter = "A"
        end_letter = get_column_letter(col_count)
        ws.auto_filter.ref = f"{start_letter}1:{end_letter}{row_count}"

    # Apply Auto-fit Column Width
    if apply_autowidth:
        for c_idx, length in max_lens.items():
            col_letter = get_column_letter(c_idx)
            # Add padding and set min/max bounds
            calculated_width = max(length + 4, 11)
            ws.column_dimensions[col_letter].width = min(calculated_width, 60)

    # Save to output stream
    wb.save(output_stream)


def format_xlsx_cell_value(val: Any) -> str:
    """Format an Excel cell value for CSV conversion and preview."""
    if val is None:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date, datetime.time)):
        if isinstance(val, datetime.datetime):
            if val.time() == datetime.time(0, 0, 0):
                return val.strftime("%Y-%m-%d")
            return val.strftime("%Y-%m-%d %H:%M:%S")
        elif isinstance(val, datetime.date):
            return val.isoformat()
        elif isinstance(val, datetime.time):
            return val.isoformat()
    if isinstance(val, bool):
        return "true" if val else "false"
    if isinstance(val, float):
        if math.isnan(val) or math.isinf(val):
            return ""
        if val.is_integer():
            return str(int(val))
        return str(val)
    return str(val)


def read_xlsx_preview(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
    max_preview_rows: int = 20,
    filename: str = "",
) -> dict:
    """Parse XLSX bytes to extract sheet names, metadata, and preview rows.

    Args:
        file_bytes: Raw bytes of uploaded XLSX file.
        sheet_name: Optional sheet name to preview. If not specified, uses active sheet.
        max_preview_rows: Number of preview rows to return (default: 20).
        filename: Original file name.

    Returns:
        Dictionary containing sheet names, active sheet, dimensions, headers, and preview rows.
    """
    validate_xlsx_content(file_bytes)

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as e:
        raise ConversionError(f"Gagal membaca file Excel (.xlsx): {str(e)}")

    sheet_names = wb.sheetnames
    if not sheet_names:
        wb.close()
        raise ConversionError("File Excel tidak memiliki lembar kerja (worksheet).")

    if sheet_name and sheet_name in sheet_names:
        target_sheet = sheet_name
    elif wb.active and wb.active.title in sheet_names:
        target_sheet = wb.active.title
    else:
        target_sheet = sheet_names[0]

    ws = wb[target_sheet]

    rows_data: List[List[str]] = []
    headers: List[str] = []
    total_rows = 0
    max_cols = 0
    empty_buffer_count = 0

    try:
        for row in ws.iter_rows(values_only=True):
            formatted_row = [format_xlsx_cell_value(c) for c in row] if row else []
            is_empty = not any(bool(c.strip()) for c in formatted_row)

            if is_empty:
                if total_rows > 0:
                    empty_buffer_count += 1
            else:
                total_rows += empty_buffer_count + 1
                empty_buffer_count = 0

                # Determine max columns
                last_idx = 0
                for idx, c in enumerate(formatted_row):
                    if c.strip():
                        last_idx = idx + 1
                if last_idx > max_cols:
                    max_cols = last_idx

                if not headers:
                    headers = formatted_row
                elif len(rows_data) < max_preview_rows:
                    rows_data.append(formatted_row)
    finally:
        wb.close()

    if total_rows == 0 or max_cols == 0:
        raise ConversionError(f"Lembar kerja '{target_sheet}' kosong atau tidak memiliki data.")

    # Normalize column lengths up to max_cols
    normalized_headers = headers[:max_cols] + [""] * max(0, max_cols - len(headers))
    normalized_rows = [
        row[:max_cols] + [""] * max(0, max_cols - len(row))
        for row in rows_data
    ]

    return {
        "success": True,
        "filename": filename,
        "sheets": sheet_names,
        "active_sheet": target_sheet,
        "total_rows": total_rows,
        "total_columns": max_cols,
        "headers": normalized_headers,
        "preview_rows": normalized_rows,
    }


def convert_xlsx_to_csv(
    file_bytes: bytes,
    output_stream: io.BufferedIOBase | io.TextIOBase,
    sheet_name: Optional[str] = None,
    delimiter: str = ",",
    encoding: str = "utf-8-sig",
) -> None:
    """Convert an Excel (.xlsx) worksheet to a delimited CSV file written to output_stream.

    Args:
        file_bytes: Raw bytes of uploaded XLSX file.
        output_stream: Writable binary or text stream.
        sheet_name: Specific worksheet name to export. Defaults to active sheet.
        delimiter: Delimiter character (',', ';', '\t', '|'). Defaults to ','.
        encoding: Output encoding ('utf-8-sig', 'utf-8', 'latin-1', 'windows-1252'). Defaults to 'utf-8-sig'.
    """
    validate_xlsx_content(file_bytes)

    if delimiter not in ALLOWED_DELIMITERS:
        delimiter = ","

    clean_enc = encoding.lower().strip() if encoding else "utf-8-sig"
    if clean_enc in ("utf8-bom", "utf-8-bom", "utf-8-sig", "utf8-sig"):
        clean_enc = "utf-8-sig"
    elif clean_enc in ("utf-8", "utf8"):
        clean_enc = "utf-8"
    elif clean_enc in ("latin-1", "latin1", "iso-8859-1"):
        clean_enc = "latin-1"
    elif clean_enc in ("windows-1252", "cp1252"):
        clean_enc = "windows-1252"
    else:
        clean_enc = "utf-8-sig"

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as e:
        raise ConversionError(f"Gagal membaca file Excel (.xlsx): {str(e)}")

    sheet_names = wb.sheetnames
    if not sheet_names:
        wb.close()
        raise ConversionError("File Excel tidak memiliki lembar kerja (worksheet).")

    if sheet_name and sheet_name in sheet_names:
        target_sheet = sheet_name
    elif wb.active and wb.active.title in sheet_names:
        target_sheet = wb.active.title
    else:
        target_sheet = sheet_names[0]

    ws = wb[target_sheet]

    # Prepare text stream wrapper if binary stream
    wrapper = None
    if isinstance(output_stream, (io.RawIOBase, io.BufferedIOBase)):
        wrapper = io.TextIOWrapper(output_stream, encoding=clean_enc, newline="", write_through=True)
        text_out = wrapper
    else:
        text_out = output_stream

    writer = csv.writer(
        text_out,
        delimiter=delimiter,
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\n",
    )

    empty_rows_buffer: List[List[str]] = []
    rows_written = 0

    try:
        for row in ws.iter_rows(values_only=True):
            formatted_row = [format_xlsx_cell_value(c) for c in row] if row else []
            is_empty = not any(bool(c.strip()) for c in formatted_row)

            if is_empty:
                if rows_written > 0:
                    empty_rows_buffer.append(formatted_row)
            else:
                for eb in empty_rows_buffer:
                    writer.writerow(eb)
                    rows_written += 1
                empty_rows_buffer.clear()

                writer.writerow(formatted_row)
                rows_written += 1
    finally:
        wb.close()
        if wrapper is not None:
            wrapper.flush()
            wrapper.detach()

    if rows_written == 0:
        raise ConversionError(f"Lembar kerja '{target_sheet}' tidak mengandung data.")
