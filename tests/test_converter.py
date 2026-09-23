"""Tests for converter engine."""

import io
import openpyxl
import pytest

from app.converter import (
    ConversionError,
    convert_csv_to_xlsx,
    convert_xlsx_to_csv,
    detect_delimiter,
    detect_encoding,
    format_xlsx_cell_value,
    parse_cell_value,
    read_csv_preview,
    read_xlsx_preview,
    sanitize_sheet_name,
    validate_file_content,
    validate_xlsx_content,
)


def test_detect_delimiter():
    csv_comma = "id,name,age\n1,Alice,30\n2,Bob,25\n"
    assert detect_delimiter(csv_comma) == ","

    csv_semi = "id;name;age\n1;Alice;30\n2;Bob;25\n"
    assert detect_delimiter(csv_semi) == ";"

    csv_tab = "id\tname\tage\n1\tAlice\t30\n2\tBob\t25\n"
    assert detect_delimiter(csv_tab) == "\t"

    csv_pipe = "id|name|age\n1|Alice|30\n2|Bob|25\n"
    assert detect_delimiter(csv_pipe) == "|"


def test_detect_encoding():
    utf8_bytes = "Halo Dunia, Ini UTF-8".encode("utf-8")
    assert detect_encoding(utf8_bytes) in ("utf-8", "ascii")

    bom_bytes = b"\xef\xbb\xbf" + "Halo Dunia".encode("utf-8")
    assert detect_encoding(bom_bytes) == "utf-8-sig"

    latin1_bytes = "nama;kota;catatan\nRené;Zürich;Spécialité\nFrançois;Nîmes;Élève\n".encode("latin-1")
    enc = detect_encoding(latin1_bytes)
    assert enc in ("latin-1", "iso-8859-1", "windows-1252", "cp1252", "cp1250")
    assert latin1_bytes.decode(enc) is not None


def test_parse_cell_value():
    assert parse_cell_value("123") == 123
    assert parse_cell_value("-456") == -456
    assert parse_cell_value("12.34") == 12.34
    assert parse_cell_value("-0.5") == -0.5

    # Leading zeroes must be preserved as strings (phone numbers / codes)
    assert parse_cell_value("08123456789") == "08123456789"
    assert parse_cell_value("00123") == "00123"
    assert parse_cell_value("0") == 0

    # Booleans
    assert parse_cell_value("true") is True
    assert parse_cell_value("TRUE") is True
    assert parse_cell_value("false") is False
    assert parse_cell_value("FALSE") is False

    # Text
    assert parse_cell_value("Hello World") == "Hello World"
    assert parse_cell_value("") == ""


def test_sanitize_sheet_name():
    assert sanitize_sheet_name("Sales_2026") == "Sales_2026"
    assert sanitize_sheet_name("Sales/Report?*:[Test]") == "Sales_Report____Test_"
    assert sanitize_sheet_name("   ") == "Sheet1"
    # Max length 31 characters
    long_name = "A" * 50
    assert len(sanitize_sheet_name(long_name)) == 31


def test_validate_file_content():
    with pytest.raises(ConversionError):
        validate_file_content(b"")

    with pytest.raises(ConversionError):
        validate_file_content(b"   \n  ")

    with pytest.raises(ConversionError):
        validate_file_content(b"PK\x03\x04\x00\x00\x00\x00")


def test_read_csv_preview():
    csv_data = "id,nama,nilai\n1,Budi,85.5\n2,Siti,90\n3,Andi,78\n".encode("utf-8")
    preview = read_csv_preview(csv_data, filename="siswa.csv")

    assert preview["success"] is True
    assert preview["filename"] == "siswa.csv"
    assert preview["detected_delimiter"] == ","
    assert preview["total_rows"] == 4
    assert preview["total_columns"] == 3
    assert preview["headers"] == ["id", "nama", "nilai"]
    assert len(preview["preview_rows"]) == 3
    assert preview["preview_rows"][0] == ["1", "Budi", "85.5"]


def test_convert_csv_to_xlsx():
    csv_data = (
        "id,nama,telepon,aktif,saldo\n"
        "1,Budi Santoso,08123456789,true,1500000.50\n"
        "2,Siti Rahma,08987654321,false,2750000\n"
    ).encode("utf-8")

    out_stream = io.BytesIO()
    convert_csv_to_xlsx(
        file_bytes=csv_data,
        output_stream=out_stream,
        sheet_name="Data Pelanggan",
        apply_autofilter=True,
        apply_autowidth=True,
    )

    out_stream.seek(0)
    wb = openpyxl.load_workbook(out_stream)
    assert "Data Pelanggan" in wb.sheetnames

    ws = wb["Data Pelanggan"]
    # Check Header
    assert ws.cell(row=1, column=1).value == "id"
    assert ws.cell(row=1, column=2).value == "nama"
    assert ws.cell(row=1, column=3).value == "telepon"

    # Check Data & Smart Typing
    assert ws.cell(row=2, column=1).value == 1  # int
    assert ws.cell(row=2, column=2).value == "Budi Santoso"
    assert ws.cell(row=2, column=3).value == "08123456789"  # preserved leading zero
    assert ws.cell(row=2, column=4).value is True  # boolean
    assert ws.cell(row=2, column=5).value == 1500000.50  # float

    # Check Auto-filter
    assert ws.auto_filter.ref is not None


def test_validate_xlsx_content():
    with pytest.raises(ConversionError, match="kosong"):
        validate_xlsx_content(b"")

    with pytest.raises(ConversionError, match="bukan merupakan format spreadsheet"):
        validate_xlsx_content(b"hello world, this is text")

    with pytest.raises(ConversionError, match="rusak atau tidak dapat dibaca"):
        validate_xlsx_content(b"PK\x03\x041234567890")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A", "B"])
    buf = io.BytesIO()
    wb.save(buf)
    validate_xlsx_content(buf.getvalue())


def test_format_xlsx_cell_value():
    import datetime
    assert format_xlsx_cell_value(None) == ""
    assert format_xlsx_cell_value(123) == "123"
    assert format_xlsx_cell_value(45.0) == "45"
    assert format_xlsx_cell_value(45.5) == "45.5"
    assert format_xlsx_cell_value(True) == "true"
    assert format_xlsx_cell_value(False) == "false"
    assert format_xlsx_cell_value(datetime.date(2026, 9, 23)) == "2026-09-23"
    assert format_xlsx_cell_value(datetime.datetime(2026, 9, 23, 15, 30, 0)) == "2026-09-23 15:30:00"
    assert format_xlsx_cell_value(datetime.datetime(2026, 9, 23, 0, 0, 0)) == "2026-09-23"


def test_read_xlsx_preview():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Karyawan"
    ws1.append(["id", "nama", "jabatan"])
    ws1.append([1, "Budi Santoso", "Manager"])
    ws1.append([2, "Siti Rahma", "Developer"])

    ws2 = wb.create_sheet(title="Departemen")
    ws2.append(["dept_id", "dept_name"])
    ws2.append(["IT", "Teknologi Informasi"])

    buf = io.BytesIO()
    wb.save(buf)
    raw_xlsx = buf.getvalue()

    # Preview default (first) sheet
    preview = read_xlsx_preview(raw_xlsx, filename="data.xlsx")
    assert preview["success"] is True
    assert preview["filename"] == "data.xlsx"
    assert preview["sheets"] == ["Karyawan", "Departemen"]
    assert preview["active_sheet"] == "Karyawan"
    assert preview["total_rows"] == 3
    assert preview["total_columns"] == 3
    assert preview["headers"] == ["id", "nama", "jabatan"]
    assert len(preview["preview_rows"]) == 2
    assert preview["preview_rows"][0] == ["1", "Budi Santoso", "Manager"]

    # Preview second sheet specifically
    preview_dept = read_xlsx_preview(raw_xlsx, sheet_name="Departemen", filename="data.xlsx")
    assert preview_dept["active_sheet"] == "Departemen"
    assert preview_dept["headers"] == ["dept_id", "dept_name"]
    assert preview_dept["total_rows"] == 2
    assert preview_dept["preview_rows"][0] == ["IT", "Teknologi Informasi"]


def test_convert_xlsx_to_csv_delimiters_and_encodings():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produk"
    ws.append(["kode", "nama_produk", "harga", "tersedia"])
    ws.append(["P01", "Kopi Robusta", 25000, True])
    ws.append(["P02", "Teh Melati", 15000.5, False])

    buf = io.BytesIO()
    wb.save(buf)
    raw_xlsx = buf.getvalue()

    # 1. Comma with default utf-8-sig
    out_sig = io.BytesIO()
    convert_xlsx_to_csv(raw_xlsx, out_sig, sheet_name="Produk", delimiter=",", encoding="utf-8-sig")
    sig_bytes = out_sig.getvalue()
    assert sig_bytes.startswith(b"\xef\xbb\xbf")
    decoded = sig_bytes.decode("utf-8-sig")
    assert "kode,nama_produk,harga,tersedia\n" in decoded
    assert "P01,Kopi Robusta,25000,true\n" in decoded
    assert "P02,Teh Melati,15000.5,false" in decoded

    # 2. Semicolon with utf-8 (no BOM)
    out_semi = io.BytesIO()
    convert_xlsx_to_csv(raw_xlsx, out_semi, sheet_name="Produk", delimiter=";", encoding="utf-8")
    semi_bytes = out_semi.getvalue()
    assert not semi_bytes.startswith(b"\xef\xbb\xbf")
    decoded_semi = semi_bytes.decode("utf-8")
    assert "kode;nama_produk;harga;tersedia\n" in decoded_semi
    assert "P01;Kopi Robusta;25000;true\n" in decoded_semi

    # 3. Tab delimiter
    out_tab = io.BytesIO()
    convert_xlsx_to_csv(raw_xlsx, out_tab, sheet_name="Produk", delimiter="\t")
    decoded_tab = out_tab.getvalue().decode("utf-8-sig")
    assert "kode\tnama_produk\tharga\ttersedia\n" in decoded_tab


def test_convert_xlsx_empty_sheet_error():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kosong"
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ConversionError, match="tidak mengandung data|kosong"):
        convert_xlsx_to_csv(buf.getvalue(), io.BytesIO())
