"""Integration tests for FastAPI endpoints."""

import io
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def test_health_check():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_preview_endpoint_success():
    csv_content = b"colA,colB,colC\nval1,val2,100\nval3,val4,200\n"
    response = client.post(
        "/api/preview",
        files={"file": ("sample.csv", csv_content, "text/csv")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["filename"] == "sample.csv"
    assert data["detected_delimiter"] == ","
    assert data["total_rows"] == 3
    assert data["total_columns"] == 3
    assert data["headers"] == ["colA", "colB", "colC"]
    assert len(data["preview_rows"]) == 2


def test_preview_endpoint_empty_file():
    response = client.post(
        "/api/preview",
        files={"file": ("empty.csv", b"", "text/csv")},
    )
    assert response.status_code == 400
    assert "File CSV kosong" in response.json()["detail"]


def test_preview_endpoint_binary_rejection():
    bad_data = b"\x00\x01\x02\x03\x04\x05\x00"
    response = client.post(
        "/api/preview",
        files={"file": ("bad.bin", bad_data, "application/octet-stream")},
    )
    assert response.status_code == 400
    assert "biner" in response.json()["detail"]


def test_convert_endpoint_success():
    csv_content = b"nama,gaji,status\nBudi,5000000,true\nAni,6500000,false\n"
    response = client.post(
        "/api/convert",
        files={"file": ("karyawan.csv", csv_content, "text/csv")},
        data={"sheet_name": "Daftar Karyawan", "apply_autofilter": "true", "apply_autowidth": "true"},
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="karyawan.xlsx"' in response.headers["content-disposition"]

    # Verify that the generated binary is indeed a readable XLSX file
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Daftar Karyawan" in wb.sheetnames
    ws = wb["Daftar Karyawan"]
    assert ws.cell(row=1, column=1).value == "nama"
    assert ws.cell(row=2, column=2).value == 5000000


def test_index_page_served():
    response = client.get("/")
    assert response.status_code == 200
    assert "csv2xlsx" in response.text


def test_preview_xlsx_endpoint_success():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Rekap"
    ws1.append(["bulan", "omset"])
    ws1.append(["Januari", 10000000])

    ws2 = wb.create_sheet("Target")
    ws2.append(["target_omset"])
    ws2.append([12000000])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    # Default sheet preview
    response = client.post(
        "/api/xlsx/preview",
        files={"file": ("laporan.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["filename"] == "laporan.xlsx"
    assert data["sheets"] == ["Rekap", "Target"]
    assert data["active_sheet"] == "Rekap"
    assert data["headers"] == ["bulan", "omset"]
    assert data["preview_rows"] == [["Januari", "10000000"]]

    # Specific sheet preview
    response_target = client.post(
        "/api/xlsx/preview",
        files={"file": ("laporan.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"sheet_name": "Target"},
    )
    assert response_target.status_code == 200
    assert response_target.json()["active_sheet"] == "Target"
    assert response_target.json()["headers"] == ["target_omset"]


def test_preview_xlsx_endpoint_invalid_file():
    response = client.post(
        "/api/xlsx/preview",
        files={"file": ("not_excel.xlsx", b"Ini bukan berkas zip excel", "application/octet-stream")},
    )
    assert response.status_code == 400
    assert "bukan merupakan format spreadsheet" in response.json()["detail"]


def test_convert_xlsx_endpoint_success():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Barang"
    ws.append(["id", "nama", "stok"])
    ws.append([101, "Meja", 25])

    buf = io.BytesIO()
    wb.save(buf)

    response = client.post(
        "/api/xlsx/convert",
        files={"file": ("stok_barang.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"delimiter": ",", "encoding": "utf-8-sig"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert 'filename="stok_barang.csv"' in response.headers["content-disposition"]
    # Verify content with utf-8-sig
    assert response.content.startswith(b"\xef\xbb\xbf")
    decoded = response.content.decode("utf-8-sig")
    assert "id,nama,stok\n101,Meja,25\n" in decoded


def test_convert_xlsx_endpoint_custom_delimiter_and_sheet():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "SheetUtama"
    ws1.append(["x", "y"])

    ws2 = wb.create_sheet("DataKhusus")
    ws2.append(["kode", "keterangan"])
    ws2.append(["A01", "Aktif"])

    buf = io.BytesIO()
    wb.save(buf)

    response = client.post(
        "/api/xlsx/convert",
        files={"file": ("proyek.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"sheet_name": "DataKhusus", "delimiter": ";", "encoding": "utf-8"},
    )
    assert response.status_code == 200
    assert 'filename="proyek_DataKhusus.csv"' in response.headers["content-disposition"]
    assert not response.content.startswith(b"\xef\xbb\xbf")
    decoded = response.content.decode("utf-8")
    assert "kode;keterangan\nA01;Aktif\n" in decoded
